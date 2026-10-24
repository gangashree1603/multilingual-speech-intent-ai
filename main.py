"""AI Project Builder backend: analyze -> generate -> download / publish to GitHub."""
import base64, io, json, os, re, secrets, uuid, zipfile
from pathlib import PurePosixPath
from typing import Literal
from urllib.parse import urlencode

import httpx
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import RedirectResponse, StreamingResponse
from pydantic import BaseModel, Field
from starlette.middleware.sessions import SessionMiddleware

load_dotenv()
E = os.environ.get
FRONTEND_URL = E("FRONTEND_URL", "http://localhost:5173")
GH = "https://api.github.com"
MAX_UPLOAD, MAX_FILES, MAX_FILE_BYTES = 5 * 1024 * 1024, 150, 300_000

app = FastAPI(title="AI Project Builder")
app.add_middleware(SessionMiddleware, secret_key=E("SESSION_SECRET") or secrets.token_hex(32),
                   same_site="lax", https_only=E("COOKIE_SECURE", "false") == "true", max_age=86400)
app.add_middleware(CORSMiddleware, allow_origins=[FRONTEND_URL], allow_credentials=True,
                   allow_methods=["GET", "POST"], allow_headers=["Content-Type"])

# Server-side state. Tokens never reach the browser (the cookie only holds a random session id).
PROJECTS: dict[str, dict] = {}   # project_id -> {sid, files:{path:content}, skipped:[]}
TOKENS: dict[str, dict] = {}     # sid -> {token, login}

# ---------------------------------------------------------------- helpers
BLOCKED_DIRS = {"node_modules", ".git", "__pycache__", "dist", "build", ".venv", "venv", ".next", ".cache", ".pytest_cache"}
SECRET_RE = re.compile(r"(ghp_[A-Za-z0-9]{36}|github_pat_\w{20,}|sk-[A-Za-z0-9_-]{20,}|AKIA[0-9A-Z]{16}|-----BEGIN [A-Z ]*PRIVATE KEY-----)")


def clean_path(p: str):
    """Return a safe relative POSIX path, or None if the file must never be included."""
    parts = PurePosixPath(str(p).replace("\\", "/").strip().lstrip("/")).parts
    if not parts or any(x in ("..", ".") or ":" in x or "\x00" in x for x in parts):
        return None
    if any(x in BLOCKED_DIRS for x in parts):
        return None
    n = parts[-1].lower()
    if n == ".env" or (n.startswith(".env.") and n != ".env.example") or n.endswith((".pem", ".key", ".pyc")) or n.startswith("id_rsa"):
        return None
    return "/".join(parts)


def sid_of(request: Request) -> str:
    return request.session.setdefault("sid", secrets.token_hex(16))


def owned(pid: str, request: Request) -> dict:
    pr = PROJECTS.get(pid)
    if not pr or pr["sid"] != sid_of(request):
        raise HTTPException(404, "Project not found.")
    return pr


def parse_json(text: str):
    a, b = text.find("{"), text.rfind("}")
    if a < 0 or b < 0:
        raise ValueError("no JSON object")
    return json.loads(text[a:b + 1])


async def ask_ai(system: str, user: str, max_tokens: int = 16000) -> str:
    prov, key, model = E("AI_PROVIDER", "anthropic").lower(), E("AI_API_KEY"), E("AI_MODEL")
    if not key or not model:
        raise HTTPException(503, "AI is not configured: set AI_API_KEY and AI_MODEL in backend/.env.")
    try:
        async with httpx.AsyncClient(timeout=600) as c:
            if prov == "anthropic":
                r = await c.post(E("AI_BASE_URL", "https://api.anthropic.com") + "/v1/messages",
                                 headers={"x-api-key": key, "anthropic-version": "2023-06-01"},
                                 json={"model": model, "max_tokens": max_tokens, "system": system,
                                       "messages": [{"role": "user", "content": user}]})
                r.raise_for_status()
                return "".join(b.get("text", "") for b in r.json()["content"])
            r = await c.post(E("AI_BASE_URL", "https://api.openai.com/v1") + "/chat/completions",
                             headers={"Authorization": f"Bearer {key}"},
                             json={"model": model, "max_tokens": max_tokens,
                                   "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}]})
            r.raise_for_status()
            return r.json()["choices"][0]["message"]["content"]
    except httpx.HTTPStatusError as e:
        raise HTTPException(502, f"AI provider returned HTTP {e.response.status_code}.")
    except httpx.HTTPError:
        raise HTTPException(502, "Could not reach the AI provider.")


async def extract_text(f: UploadFile) -> str:
    data = await f.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "File is larger than 5 MB.")
    ext = os.path.splitext(f.filename or "")[1].lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
        if ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs)
        if ext == ".txt":
            return data.decode("utf-8", "replace")
    except Exception:
        raise HTTPException(422, "Could not read that document. Is it corrupted or password-protected?")
    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        raise HTTPException(415, "Text extraction from images needs OCR, which is not enabled. Paste the text instead.")
    raise HTTPException(415, "Unsupported file type. Use PDF, DOCX or TXT.")

# ---------------------------------------------------------------- projects
@app.get("/api/session")
async def session(request: Request):
    t = TOKENS.get(sid_of(request))
    return {"github": t["login"] if t else None, "ai_configured": bool(E("AI_API_KEY") and E("AI_MODEL")),
            "github_configured": bool(E("GITHUB_CLIENT_ID") and E("GITHUB_CLIENT_SECRET"))}


@app.post("/api/projects/analyze")
async def analyze(request: Request, text: str = Form(""), file: UploadFile | None = File(None)):
    sid_of(request)
    spec = text.strip()
    if file and file.filename:
        spec = (spec + "\n\n" + (await extract_text(file)).strip()).strip()
    if len(spec) < 10:
        raise HTTPException(422, "Enter a project topic or upload a document with at least a few words.")
    spec = spec[:20000]
    raw = await ask_ai("You analyze software project specs. Reply with JSON only.",
                       'Return {"title":str,"summary":str,"features":[str],"suggested_stack":str,"database":str} for this spec:\n\n' + spec, 1500)
    try:
        analysis = parse_json(raw)
    except Exception:
        raise HTTPException(502, "The AI returned an unreadable analysis. Try again.")
    return {"spec": spec, "analysis": analysis}


class GenReq(BaseModel):
    spec: str = Field(min_length=10, max_length=20000)
    stack: str = Field(default="React + FastAPI", max_length=100)
    analysis: dict | None = None


GEN_SYSTEM = (
    "You are a senior engineer generating a COMPLETE, runnable project. Reply with JSON only, no prose or fences: "
    '{"files":[{"path":"relative/path","content":"full file text"}]}. Include frontend, backend, database config, '
    "API integration where needed, README.md (setup + run), .gitignore, dependency files with pinned versions, and "
    ".env.example with placeholder values only. Never include real secrets, .env, node_modules or build output. "
    "Use at most 60 files; every file must be complete (no TODO placeholders).")


@app.post("/api/projects/generate")
async def generate(body: GenReq, request: Request):
    sid = sid_of(request)
    raw = await ask_ai(GEN_SYSTEM, f"Preferred stack: {body.stack}\nAnalysis: {json.dumps(body.analysis or {})}\n\nSpecification:\n{body.spec}")
    try:
        items = parse_json(raw)["files"]
        assert isinstance(items, list)
    except Exception:
        raise HTTPException(502, "The AI response was not valid project JSON (it may have been cut off). Try a smaller scope.")
    files, skipped = {}, []
    for it in items[:MAX_FILES]:
        path, content = clean_path(it.get("path", "")), it.get("content", "")
        if not path or not isinstance(content, str):
            skipped.append(str(it.get("path")) + " (unsafe path)"); continue
        if len(content.encode()) > MAX_FILE_BYTES:
            skipped.append(path + " (too large)"); continue
        if SECRET_RE.search(content):
            skipped.append(path + " (looked like it contained a secret)"); continue
        files[path] = content
    if not files:
        raise HTTPException(502, "No usable files were generated.")
    pid = uuid.uuid4().hex
    PROJECTS[pid] = {"sid": sid, "files": files, "skipped": skipped}
    return {"project_id": pid, "file_count": len(files), "skipped": skipped}


@app.get("/api/projects/{project_id}/files")
async def list_files(project_id: str, request: Request):
    pr = owned(project_id, request)
    return {"project_id": project_id, "skipped": pr["skipped"],
            "files": [{"path": p, "size": len(c.encode()), "content": c} for p, c in sorted(pr["files"].items())]}


@app.get("/api/projects/{project_id}/download")
async def download(project_id: str, request: Request):
    pr = owned(project_id, request)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for p, c in pr["files"].items():
            z.writestr(p, c)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/zip",
                             headers={"Content-Disposition": 'attachment; filename="project.zip"'})

# ---------------------------------------------------------------- GitHub
@app.get("/api/github/login")
async def gh_login(request: Request):
    if not (E("GITHUB_CLIENT_ID") and E("GITHUB_CLIENT_SECRET")):
        raise HTTPException(503, "GitHub OAuth is not configured on the server.")
    state = request.session["gh_state"] = secrets.token_urlsafe(24)
    q = urlencode({"client_id": E("GITHUB_CLIENT_ID"), "redirect_uri": E("GITHUB_REDIRECT_URI", ""),
                   "scope": E("GITHUB_SCOPE", "repo"), "state": state})
    return RedirectResponse("https://github.com/login/oauth/authorize?" + q)


@app.get("/api/github/callback")
async def gh_callback(request: Request, code: str = "", state: str = "", error: str = ""):
    expected = request.session.pop("gh_state", None)
    if error or not code or not expected or not secrets.compare_digest(expected, state):
        return RedirectResponse(FRONTEND_URL + "/?github=error")
    async with httpx.AsyncClient(timeout=30) as c:
        r = await c.post("https://github.com/login/oauth/access_token", headers={"Accept": "application/json"},
                         data={"client_id": E("GITHUB_CLIENT_ID"), "client_secret": E("GITHUB_CLIENT_SECRET"),
                               "code": code, "redirect_uri": E("GITHUB_REDIRECT_URI", "")})
        tok = r.json().get("access_token") if r.status_code == 200 else None
        if not tok:
            return RedirectResponse(FRONTEND_URL + "/?github=error")
        u = await gh(c, "GET", "/user", tok)
    TOKENS[sid_of(request)] = {"token": tok, "login": u.json().get("login", "")}
    return RedirectResponse(FRONTEND_URL + "/?github=connected")


@app.post("/api/github/logout")
async def gh_logout(request: Request):
    TOKENS.pop(sid_of(request), None)
    return {"ok": True}


def token_of(request: Request) -> str:
    t = TOKENS.get(sid_of(request))
    if not t:
        raise HTTPException(401, "Connect your GitHub account first.")
    return t["token"]


async def gh(c: httpx.AsyncClient, method: str, path: str, tok: str, **kw):
    try:
        r = await c.request(method, GH + path, headers={"Authorization": f"Bearer {tok}", "Accept": "application/vnd.github+json",
                                                         "X-GitHub-Api-Version": "2022-11-28"}, **kw)
    except httpx.HTTPError:
        raise HTTPException(502, "Could not reach GitHub.")
    if r.status_code == 401:
        raise HTTPException(401, "GitHub authorization expired or was revoked. Reconnect GitHub.")
    if r.status_code in (403, 429) and (r.headers.get("x-ratelimit-remaining") == "0" or "retry-after" in r.headers or "rate limit" in r.text.lower()):
        raise HTTPException(429, "GitHub rate limit reached. Wait a few minutes and try again.")
    return r


@app.get("/api/github/repositories")
async def repos(request: Request):
    tok = token_of(request)
    async with httpx.AsyncClient(timeout=30) as c:
        r = await gh(c, "GET", "/user/repos?per_page=100&sort=updated&affiliation=owner,collaborator,organization_member", tok)
    if r.status_code != 200:
        raise HTTPException(502, f"GitHub returned HTTP {r.status_code}.")
    return [{"full_name": x["full_name"], "private": x["private"], "default_branch": x["default_branch"]}
            for x in r.json() if x.get("permissions", {}).get("push")]


class PublishReq(BaseModel):
    project_id: str
    mode: Literal["new", "existing"]
    repo: str = Field(min_length=1, max_length=140)       # name (new) or owner/name (existing)
    description: str = Field(default="", max_length=350)
    private: bool = True
    branch: str = Field(default="main", max_length=100)
    overwrite: bool = False
    confirm: bool = False


@app.post("/api/github/publish")
async def publish(body: PublishReq, request: Request):
    tok, pr = token_of(request), owned(body.project_id, request)
    if not body.confirm:
        raise HTTPException(400, "Confirm the file list before publishing.")
    if not re.fullmatch(r"[A-Za-z0-9._/-]{1,100}", body.branch) or ".." in body.branch or body.branch.startswith("/"):
        raise HTTPException(422, "Invalid branch name.")
    pattern = r"[A-Za-z0-9._-]{1,100}" if body.mode == "new" else r"[A-Za-z0-9._-]+/[A-Za-z0-9._-]+"
    if not re.fullmatch(pattern, body.repo):
        raise HTTPException(422, "Invalid repository name.")
    files = {p: c for p, c in pr["files"].items() if clean_path(p)}
    async with httpx.AsyncClient(timeout=60) as c:
        if body.mode == "new":
            r = await gh(c, "POST", "/user/repos", tok, json={"name": body.repo, "description": body.description,
                                                              "private": body.private, "auto_init": True})
            if r.status_code == 422:
                raise HTTPException(409, "A repository with that name already exists on your account (or the name is invalid).")
            if r.status_code != 201:
                raise HTTPException(502, f"GitHub could not create the repository (HTTP {r.status_code}).")
        else:
            r = await gh(c, "GET", f"/repos/{body.repo}", tok)
            if r.status_code == 404:
                raise HTTPException(404, "Repository not found or not accessible.")
            if r.status_code != 200 or not r.json().get("permissions", {}).get("push"):
                raise HTTPException(403, "You do not have push access to that repository.")
        repo = r.json(); full = repo["full_name"]

        ref = await gh(c, "GET", f"/repos/{full}/git/ref/heads/{body.branch}", tok)
        new_branch = ref.status_code == 404
        if new_branch:
            ref = await gh(c, "GET", f"/repos/{full}/git/ref/heads/{repo['default_branch']}", tok)
        if ref.status_code != 200:
            raise HTTPException(409, "The repository is empty. Create an initial commit (e.g. a README) on GitHub first, then retry.")
        parent = ref.json()["object"]["sha"]
        cm = await gh(c, "GET", f"/repos/{full}/git/commits/{parent}", tok)
        base_tree = cm.json()["tree"]["sha"]

        if body.mode == "existing" and not new_branch:
            t = await gh(c, "GET", f"/repos/{full}/git/trees/{base_tree}?recursive=1", tok)
            existing = {x["path"] for x in t.json().get("tree", []) if x["type"] == "blob"}
            conflicts = sorted(existing & set(files))
            if conflicts and not body.overwrite:
                raise HTTPException(409, {"message": f"{len(conflicts)} file(s) already exist. Tick 'overwrite' to replace them.", "conflicts": conflicts})

        entries = []
        for p, content in files.items():
            b = await gh(c, "POST", f"/repos/{full}/git/blobs", tok,
                         json={"content": base64.b64encode(content.encode()).decode(), "encoding": "base64"})
            if b.status_code != 201:
                raise HTTPException(502, f"GitHub rejected {p} (HTTP {b.status_code}). Nothing was committed.")
            entries.append({"path": p, "mode": "100644", "type": "blob", "sha": b.json()["sha"]})
        t = await gh(c, "POST", f"/repos/{full}/git/trees", tok, json={"base_tree": base_tree, "tree": entries})
        if t.status_code != 201:
            raise HTTPException(502, f"GitHub could not build the file tree (HTTP {t.status_code}).")
        cm = await gh(c, "POST", f"/repos/{full}/git/commits", tok,
                      json={"message": "Add project generated by AI Project Builder", "tree": t.json()["sha"], "parents": [parent]})
        if cm.status_code != 201:
            raise HTTPException(502, f"GitHub could not create the commit (HTTP {cm.status_code}).")
        sha = cm.json()["sha"]
        if new_branch:
            u = await gh(c, "POST", f"/repos/{full}/git/refs", tok, json={"ref": f"refs/heads/{body.branch}", "sha": sha})
            ok = u.status_code == 201
        else:
            u = await gh(c, "PATCH", f"/repos/{full}/git/refs/heads/{body.branch}", tok, json={"sha": sha, "force": False})
            ok = u.status_code == 200
        if not ok:
            if u.status_code == 422:
                raise HTTPException(409, "The branch changed while publishing. Retry to build on the latest commit.")
            raise HTTPException(502, f"GitHub did not accept the update (HTTP {u.status_code}). Nothing was published.")
    return {"published": True, "repo_url": repo["html_url"], "commit_url": cm.json()["html_url"], "branch": body.branch, "files": len(entries)}
